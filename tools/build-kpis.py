"""Generate config/kpis.json from the KPI catalogue Excel.

Usage: python tools/build-kpis.py path/to/CIO_Toolkit_KPI_Catalogue.xlsx [out.json]
Needs: pip install openpyxl
"""
import json
import sys
from pathlib import Path

from openpyxl import load_workbook

# Excel column -> JSON field. The Excel column "Dashboard view" holds the reporting cadence.
FIELDS = {
    "ID": "id",
    "Domain": "domain",
    "KPI": "name",
    "Definition / formula": "definition",
    "Typical data source": "dataSource",
    "Refresh": "refresh",
    "Suggested owner": "owner",
    "Dashboard view": "cadence",
    "Integration complexity": "effort",
    "AI value": "aiValue",
    "AI role": "aiRole",
    "How AI helps (or why not)": "aiHelp",
    "Risk / caution": "risk",
    "MVP candidate": "mvp",
    "Comments from Leaders": "leaderComments",
}


def clean(value):
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def main(xlsx, out):
    ws = load_workbook(xlsx, data_only=True).active
    rows = list(ws.iter_rows(values_only=True))
    header_at = next(i for i, r in enumerate(rows) if r and "ID" in r and "KPI" in r)
    header = [clean(h) for h in rows[header_at]]
    missing = [c for c in FIELDS if c not in header]
    if missing:
        sys.exit("Missing columns in the Excel: " + ", ".join(missing))
    col = {name: header.index(name) for name in FIELDS}

    kpis = []
    for r in rows[header_at + 1:]:
        if not r or not clean(r[col["ID"]]):
            continue
        item = {}
        for name, field in FIELDS.items():
            value = clean(r[col[name]])
            if field == "mvp":
                value = bool(value) and value.upper().startswith("Y")
            item[field] = value
        kpis.append(item)

    Path(out).write_text(json.dumps(kpis, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print(f"Wrote {len(kpis)} KPIs to {out}; MVP candidates: {sum(k['mvp'] for k in kpis)}")


if __name__ == "__main__":
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    main(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else "config/kpis.json")
